from io import BytesIO
from pathlib import Path
import html
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt
from fpdf import FPDF

ASSETS = Path(__file__).resolve().parents[2] / "assets"
LOGO = ASSETS / "logo.png"

def sanitize_text(text: str) -> str:
    text = text.replace("\u2018", "'").replace("\u2019", "'")
    text = text.replace("\u201c", '"').replace("\u201d", '"')
    text = text.replace("\u2013", "-").replace("\u2014", "-").replace("\u00a0", " ")
    return re.sub(r"[^\S\r\n]+", " ", text).strip()

def format_html_preview(text: str) -> str:
    safe = html.escape(sanitize_text(text))
    safe = re.sub(r"(?m)^#{1,6}\s+(.+)$", r"<h3>\1</h3>", safe)
    safe = re.sub(r"(?m)^\*\*(.+?)\*\*$", r"<strong>\1</strong>", safe)
    safe = re.sub(r"(?m)^\s*[-•]\s+(.+)$", r"<li>\1</li>", safe)
    safe = re.sub(r"(<li>.*?</li>\n?)+", lambda m: f"<ul>{m.group(0)}</ul>", safe)
    return safe.replace("\n", "<br>")

def _lines(text: str) -> list[str]:
    return [line.strip() for line in sanitize_text(text).splitlines() if line.strip()]

def format_docx(text: str, doc_type: str, terms: str | None = None) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(1.25))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for line in _lines(text):
        if re.match(r"^(\d+[\.\)]|[A-Z][A-Z\s&/-]{3,}:|#{1,6}\s)", line):
            p = doc.add_paragraph()
            r = p.add_run(re.sub(r"^#{1,6}\s*", "", line))
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
        else:
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(7)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)

    if terms:
        bullets = [x.strip() for x in terms.split(";") if x.strip()]
        if bullets:
            doc.add_paragraph().add_run("KEY TERMS").bold = True
            table = doc.add_table(rows=1, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "No."
            table.rows[0].cells[1].text = "Term"
            for i, item in enumerate(bullets, 1):
                cells = table.add_row().cells
                cells[0].text = str(i)
                cells[1].text = item

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase - AI-generated template. Review with a qualified legal professional before use.")

    out = BytesIO()
    doc.save(out)
    return out.getvalue()

class BrandedPDF(FPDF):
    def __init__(self, doc_type: str):
        super().__init__()
        self.doc_type = doc_type
        self.set_auto_page_break(auto=True, margin=18)

    def header(self):
        if LOGO.exists():
            self.image(str(LOGO), x=95, y=8, w=20)
            self.ln(16)
        self.set_font("Times", "B", 12)
        self.cell(0, 8, self.doc_type.upper(), align="C")
        self.ln(8)

    def footer(self):
        self.set_y(-15)
        self.set_font("Times", "I", 8)
        self.cell(0, 8, "LegalEase - AI-generated template; review before use.", align="C")

def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = BrandedPDF(doc_type)
    pdf.set_margins(18, 25, 18)
    pdf.add_page()
    for line in _lines(text):
        if re.match(r"^(\d+[\.\)]|[A-Z][A-Z\s&/-]{3,}:|#{1,6}\s)", line):
            pdf.set_font("Times", "B", 11)
            pdf.multi_cell(0, 7, re.sub(r"^#{1,6}\s*", "", line))
            pdf.ln(1)
        elif re.match(r"^[-•]\s+", line):
            pdf.set_font("Times", "", 11)
            pdf.multi_cell(0, 6, "- " + re.sub(r"^[-•]\s+", "", line))
        else:
            pdf.set_font("Times", "", 11)
            pdf.multi_cell(0, 6, line)
            pdf.ln(1)
    return bytes(pdf.output())

def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")
