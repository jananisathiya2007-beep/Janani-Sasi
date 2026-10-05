import streamlit as st
import requests
import os
from dotenv import load_dotenv


# Load .env
load_dotenv()


# -----------------------------------
# Page configuration
# -----------------------------------

st.set_page_config(
    page_title="LegalEase AI",
    page_icon="⚖️",
    layout="wide"
)


# -----------------------------------
# Custom CSS
# -----------------------------------

st.markdown(
    """
    <style>

    .main-title {
        text-align: center;
        font-size: 45px;
        font-weight: bold;
        margin-bottom: 5px;
    }

    .subtitle {
        text-align: center;
        font-size: 18px;
        color: #666666;
        margin-bottom: 30px;
    }

    .result-box {
        background-color: #f8f9fa;
        padding: 20px;
        border-radius: 10px;
        border: 1px solid #dddddd;
    }

    </style>
    """,
    unsafe_allow_html=True
)


# -----------------------------------
# Title
# -----------------------------------

st.markdown(
    '<div class="main-title">⚖️ LegalEase AI</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="subtitle">AI-Powered Legal Document Generation System</div>',
    unsafe_allow_html=True
)


st.divider()


# -----------------------------------
# Backend URL
# -----------------------------------

BACKEND_URL = os.getenv(
    "BACKEND_URL",
    "http://127.0.0.1:8000"
)


# -----------------------------------
# Input section
# -----------------------------------

st.header("📄 Create Legal Document")


document_type = st.selectbox(
    "Document Type",
    [
        "Rental Agreement",
        "Non-Disclosure Agreement (NDA)",
        "Employment Agreement",
        "Business Agreement",
        "Service Agreement",
        "Partnership Agreement",
        "Sale Agreement",
        "General Legal Agreement"
    ]
)


parties = st.text_area(
    "Parties",
    placeholder=(
        "Example:\n"
        "Party 1: ABC Technologies Pvt Ltd\n"
        "Party 2: John Smith"
    ),
    height=120
)


effective_date = st.text_input(
    "Effective Date",
    placeholder="Example: 5 October 2026"
)


terms = st.text_area(
    "Terms and Conditions",
    placeholder=(
        "Example:\n"
        "Monthly payment: Rs. 20,000\n"
        "Agreement duration: 12 months\n"
        "Either party can terminate with 30 days notice."
    ),
    height=180
)


# -----------------------------------
# Generate button
# -----------------------------------

generate_button = st.button(
    "🚀 Generate Legal Document",
    use_container_width=True
)


# -----------------------------------
# Generate document
# -----------------------------------

if generate_button:

    if not parties.strip():
        st.warning("Please enter the parties.")

    elif not effective_date.strip():
        st.warning("Please enter the effective date.")

    elif not terms.strip():
        st.warning("Please enter the terms and conditions.")

    else:

        request_data = {
            "document_type": document_type,
            "parties": parties,
            "terms": terms,
            "effective_date": effective_date
        }

        with st.spinner("🤖 Gemini AI is generating your legal document..."):

            try:

                response = requests.post(
                    f"{BACKEND_URL}/generate",
                    json=request_data,
                    timeout=120
                )


                if response.status_code == 200:

                    result = response.json()

                    st.session_state["document"] = result["document"]

                    st.success(
                        "✅ Legal document generated successfully!"
                    )

                else:

                    try:
                        error_message = response.json().get(
                            "detail",
                            "Unknown backend error."
                        )
                    except Exception:
                        error_message = response.text

                    st.error(
                        f"Backend Error: {error_message}"
                    )

            except requests.exceptions.ConnectionError:

                st.error(
                    "❌ Cannot connect to the backend. "
                    "Make sure FastAPI is running."
                )

            except Exception as error:

                st.error(
                    f"❌ Error: {error}"
                )


# -----------------------------------
# Display generated document
# -----------------------------------

if "document" in st.session_state:

    st.divider()

    st.header("📑 Generated Legal Document")


    # Editable document
    edited_document = st.text_area(
        "Edit your document",
        value=st.session_state["document"],
        height=600
    )


    # Save edits
    st.session_state["document"] = edited_document


    st.subheader("👀 Document Preview")


    st.markdown(
        '<div class="result-box">',
        unsafe_allow_html=True
    )

    st.markdown(
        edited_document
    )

    st.markdown(
        "</div>",
        unsafe_allow_html=True
    )


    st.divider()


    # -----------------------------------
    # Download options
    # -----------------------------------

    st.subheader("⬇️ Download Document")


    col1, col2, col3 = st.columns(3)


    # TXT
    with col1:

        st.download_button(
            label="📄 Download TXT",
            data=edited_document,
            file_name="LegalEase_Document.txt",
            mime="text/plain",
            use_container_width=True
        )


    # DOCX
    with col2:

        try:

            from docx import Document

            doc = Document()

            doc.add_heading(
                document_type,
                0
            )

            for paragraph in edited_document.split("\n"):

                if paragraph.strip():

                    doc.add_paragraph(
                        paragraph
                    )

            docx_path = "LegalEase_Document.docx"

            doc.save(docx_path)

            with open(
                docx_path,
                "rb"
            ) as file:

                docx_data = file.read()


            st.download_button(
                label="📝 Download DOCX",
                data=docx_data,
                file_name="LegalEase_Document.docx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                use_container_width=True
            )

        except Exception as error:

            st.error(
                f"DOCX creation error: {error}"
            )


    # PDF
    with col3:

        try:

            from fpdf import FPDF


            pdf = FPDF()

            pdf.set_auto_page_break(
                auto=True,
                margin=15
            )

            pdf.add_page()

            pdf.set_font(
                "Arial",
                size=16
            )

            pdf.multi_cell(
                0,
                10,
                document_type,
                align="C"
            )

            pdf.ln(5)

            pdf.set_font(
                "Arial",
                size=11
            )

            safe_text = (
                edited_document
                .replace("₹", "Rs.")
                .replace("–", "-")
                .replace("—", "-")
            )

            pdf.multi_cell(
                0,
                7,
                safe_text
            )

            pdf_data = bytes(
                pdf.output()
            )


            st.download_button(
                label="📕 Download PDF",
                data=pdf_data,
                file_name="LegalEase_Document.pdf",
                mime="application/pdf",
                use_container_width=True
            )

        except Exception as error:

            st.error(
                f"PDF creation error: {error}"
            )


# -----------------------------------
# Footer
# -----------------------------------

st.divider()

st.caption(
    "⚖️ LegalEase AI | AI-powered legal document drafting system"
)

st.caption(
    "This application generates drafts for educational and "
    "informational purposes. Please consult a qualified legal "
    "professional before using any document for legal purposes."
)